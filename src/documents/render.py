"""ATS-safe CV rendering to PDF and DOCX.

Layout rules: single column, no tables, no text boxes, no images, standard
section headings, real text (not outlines), generous margins. That is what
resume parsers expect.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.colors import HexColor  # type: ignore[import-untyped]
from reportlab.lib.enums import TA_LEFT  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
from reportlab.lib.styles import ParagraphStyle  # type: ignore[import-untyped]
from reportlab.lib.units import mm  # type: ignore[import-untyped]
from reportlab.platypus import (  # type: ignore[import-untyped]
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from src.observability.logging import get_logger

logger = get_logger(__name__)

_INK = HexColor("#111111")
_MUTED = HexColor("#444444")
_RULE = HexColor("#888888")


@dataclass(frozen=True)
class RoleBlock:
    title: str
    org: str
    meta: str | None
    bullets: tuple[str, ...]


@dataclass(frozen=True)
class ProjectBlock:
    name: str
    meta: str | None
    description: str
    bullets: tuple[str, ...]


@dataclass(frozen=True)
class CVData:
    name: str
    headline: str
    contact: tuple[str, ...]
    summary: str
    roles: tuple[RoleBlock, ...]
    projects: tuple[ProjectBlock, ...]
    skills: tuple[tuple[str, str], ...]
    education: tuple[str, ...]
    certifications: tuple[str, ...] = ()


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #


def _pdf_styles() -> dict[str, ParagraphStyle]:
    return {
        "name": ParagraphStyle(
            "cv-name",
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=21,
            textColor=_INK,
            spaceAfter=2,
            alignment=TA_LEFT,
        ),
        "headline": ParagraphStyle(
            "cv-headline",
            fontName="Helvetica",
            fontSize=10.5,
            leading=14,
            textColor=_MUTED,
            spaceAfter=2,
        ),
        "contact": ParagraphStyle(
            "cv-contact",
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=_MUTED,
        ),
        "section": ParagraphStyle(
            "cv-section",
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=12,
            textColor=_INK,
            spaceBefore=8,
            spaceAfter=0,
        ),
        "body": ParagraphStyle(
            "cv-body",
            fontName="Helvetica",
            fontSize=9.3,
            leading=12.0,
            textColor=_INK,
            spaceBefore=3,
        ),
        "bullet": ParagraphStyle(
            "cv-bullet",
            fontName="Helvetica",
            fontSize=9.3,
            leading=12.0,
            textColor=_INK,
            leftIndent=10,
            bulletIndent=0,
            spaceBefore=1,
        ),
        "detail": ParagraphStyle(
            "cv-detail",
            fontName="Helvetica-Oblique",
            fontSize=9,
            leading=12,
            textColor=_MUTED,
            spaceBefore=1,
        ),
    }


def _section(title: str, styles: dict[str, ParagraphStyle]) -> list[Any]:
    return [
        Paragraph(escape(title.upper()), styles["section"]),
        HRFlowable(
            width="100%",
            thickness=0.5,
            color=_RULE,
            spaceBefore=1,
            spaceAfter=4,
        ),
    ]


def render_pdf(cv: CVData, path: Path) -> Path:
    """Write an ATS-friendly PDF."""
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = _pdf_styles()
    story: list[Any] = []

    story.append(Paragraph(escape(cv.name), styles["name"]))
    if cv.headline:
        story.append(Paragraph(escape(cv.headline), styles["headline"]))
    if cv.contact:
        story.append(Paragraph(escape("  |  ".join(cv.contact)), styles["contact"]))

    if cv.summary:
        story.extend(_section("Professional Summary", styles))
        story.append(Paragraph(escape(cv.summary), styles["body"]))

    if cv.roles:
        story.extend(_section("Experience", styles))
        for role in cv.roles:
            story.append(
                Paragraph(
                    f"<b>{escape(role.title)}</b> — {escape(role.org)}",
                    styles["body"],
                )
            )
            if role.meta:
                story.append(Paragraph(escape(role.meta), styles["detail"]))
            for bullet in role.bullets:
                story.append(
                    Paragraph(escape(bullet), styles["bullet"], bulletText="•")
                )

    if cv.projects:
        story.extend(_section("Selected Projects", styles))
        for project in cv.projects:
            story.append(Paragraph(f"<b>{escape(project.name)}</b>", styles["body"]))
            if project.meta:
                story.append(Paragraph(escape(project.meta), styles["detail"]))
            if project.description:
                story.append(Paragraph(escape(project.description), styles["bullet"]))
            for bullet in project.bullets:
                story.append(
                    Paragraph(escape(bullet), styles["bullet"], bulletText="•")
                )

    if cv.skills:
        story.extend(_section("Technical Skills", styles))
        for group, items in cv.skills:
            story.append(
                Paragraph(
                    f"<b>{escape(group)}:</b> {escape(items)}",
                    styles["bullet"],
                )
            )

    if cv.education:
        story.extend(_section("Education", styles))
        for line in cv.education:
            story.append(Paragraph(escape(line), styles["bullet"], bulletText="•"))

    if cv.certifications:
        story.extend(_section("Certifications", styles))
        for line in cv.certifications:
            story.append(Paragraph(escape(line), styles["bullet"], bulletText="•"))

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"{cv.name} — CV",
        author=cv.name,
        subject=cv.headline or "Curriculum Vitae",
        creator="job-agent",
    )
    doc.build(story)
    logger.info("cv.pdf.rendered", path=str(path), bytes=path.stat().st_size)
    return path


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #


def _docx_setup(document: DocxDocument) -> None:
    section = document.sections[0]
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)


def _docx_paragraph(
    document: DocxDocument,
    text: str,
    *,
    size: float = 10,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
    space_before: float = 0,
    space_after: float = 2,
) -> Any:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return paragraph


def _docx_section(document: DocxDocument, title: str) -> None:
    paragraph = _docx_paragraph(
        document, title.upper(), size=9.5, bold=True, space_before=9, space_after=2
    )
    # Bottom border on the heading only; parsers still see a plain paragraph.
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "888888")
    borders.append(bottom)
    p_pr.append(borders)


def _docx_bullet(document: DocxDocument, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    run.font.size = Pt(10)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.left_indent = Cm(0.6)


def render_docx(cv: CVData, path: Path) -> Path:
    """Write an ATS-friendly DOCX."""
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _docx_setup(document)

    _docx_paragraph(document, cv.name, size=17, bold=True, space_after=2)
    if cv.headline:
        _docx_paragraph(document, cv.headline, size=10.5, color=RGBColor(0x44, 0x44, 0x44))
    if cv.contact:
        _docx_paragraph(
            document,
            "  |  ".join(cv.contact),
            size=9,
            color=RGBColor(0x44, 0x44, 0x44),
            space_after=4,
        )

    if cv.summary:
        _docx_section(document, "Professional Summary")
        _docx_paragraph(document, cv.summary, size=10)

    if cv.roles:
        _docx_section(document, "Experience")
        for role in cv.roles:
            _docx_paragraph(
                document, f"{role.title} — {role.org}", size=10, bold=True, space_before=4
            )
            if role.meta:
                _docx_paragraph(
                    document,
                    role.meta,
                    size=9,
                    italic=True,
                    color=RGBColor(0x44, 0x44, 0x44),
                )
            for bullet in role.bullets:
                _docx_bullet(document, bullet)

    if cv.projects:
        _docx_section(document, "Selected Projects")
        for project in cv.projects:
            _docx_paragraph(
                document, project.name, size=10, bold=True, space_before=4
            )
            if project.meta:
                _docx_paragraph(
                    document,
                    project.meta,
                    size=9,
                    italic=True,
                    color=RGBColor(0x44, 0x44, 0x44),
                )
            if project.description:
                _docx_paragraph(document, project.description, size=10)
            for bullet in project.bullets:
                _docx_bullet(document, bullet)

    if cv.skills:
        _docx_section(document, "Technical Skills")
        for group, items in cv.skills:
            paragraph = document.add_paragraph()
            label = paragraph.add_run(f"{group}: ")
            label.bold = True
            label.font.size = Pt(10)
            value = paragraph.add_run(items)
            value.font.size = Pt(10)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)

    if cv.education:
        _docx_section(document, "Education")
        for line in cv.education:
            _docx_bullet(document, line)

    if cv.certifications:
        _docx_section(document, "Certifications")
        for line in cv.certifications:
            _docx_bullet(document, line)

    document.core_properties.title = f"{cv.name} — CV"
    document.core_properties.author = cv.name
    document.core_properties.subject = cv.headline or "Curriculum Vitae"
    document.save(str(path))
    logger.info("cv.docx.rendered", path=str(path), bytes=path.stat().st_size)
    return path


__all__ = [
    "CVData",
    "ProjectBlock",
    "RoleBlock",
    "render_docx",
    "render_pdf",
]
