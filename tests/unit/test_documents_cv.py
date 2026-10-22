"""Tailored CV generation: KB loading, grounding, selection, rendering."""

from __future__ import annotations

import os
import uuid
from pathlib import Path
from types import SimpleNamespace

import pytest

os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://x:x@localhost/x")
os.environ.setdefault("REDIS_URL", "redis://localhost:6379/0")
os.environ.setdefault("TELEGRAM_BOT_TOKEN", "t")
os.environ.setdefault("TELEGRAM_CHAT_ID", "0")
os.environ.setdefault("ENCRYPTION_KEY", "QUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUE=")

from src.documents.generator import build_cv_data, filename_base, generate_cv
from src.documents.grounding import is_grounded, unsupported_terms
from src.documents.kb import load_kb
from src.documents.tailoring import (
    CVPlan,
    build_plan,
    deterministic_plan,
    sanitise_plan,
    selected_bullets,
    selected_projects,
)


def make_job(**overrides) -> SimpleNamespace:
    defaults = {
        "id": uuid.uuid4(),
        "title": "Senior Backend Engineer (Node.js)",
        "company": "Acme Payments",
        "description": (
            "Build Node.js and PostgreSQL payment services on AWS. React dashboard "
            "work. Security background a plus. Kafka event streaming."
        ),
        "technologies": ["Node.js", "PostgreSQL", "AWS", "React", "Kafka"],
        "seniority": "senior",
        "employment_type": "full-time",
        "location": "Remote (EU)",
        "remote": True,
    }
    defaults.update(overrides)
    return SimpleNamespace(**defaults)


@pytest.fixture(scope="module")
def kb():
    return load_kb()


def test_kb_loads_core_sections(kb):
    assert kb.name
    assert kb.email
    assert kb.experience, "employment history should load"
    assert any(entry.bullets for entry in kb.experience)
    assert len(kb.projects) >= 3
    assert kb.skills
    assert kb.education


def test_deterministic_plan_references_only_kb_material(kb):
    plan = deterministic_plan(make_job(), kb)
    pool = {bid for bid, _ in kb.bullet_pool()}
    project_ids = {project.id for project in kb.projects}
    skill_names = {name for names in kb.skills.values() for name in names}

    assert set(plan.experience_bullet_ids) <= pool
    assert set(plan.project_ids) <= project_ids
    assert set(plan.skill_names) <= skill_names
    assert len(plan.experience_bullet_ids) >= 6
    assert len(plan.project_ids) >= 2
    assert is_grounded(plan.summary, kb)


def test_plan_drops_unknown_ids_and_keeps_minimums(kb):
    plan = CVPlan(
        headline="Principal Cloud Architect",
        summary="I have 12 years of Kubernetes and Salesforce experience.",
        experience_bullet_ids=["not-a-real-bullet"],
        project_ids=["nonexistent-project"],
        skill_names=["Imaginary Skill"],
    )
    cleaned = sanitise_plan(plan, make_job(), kb)

    assert set(cleaned.experience_bullet_ids) <= {bid for bid, _ in kb.bullet_pool()}
    assert set(cleaned.project_ids) <= {project.id for project in kb.projects}
    assert "Imaginary Skill" not in cleaned.skill_names
    assert is_grounded(cleaned.headline, kb)
    assert is_grounded(cleaned.summary, kb)


def test_grounding_flags_invented_claims(kb):
    assert unsupported_terms("Expert in Salesforce and COBOL.", kb)
    assert unsupported_terms("Principal engineer leading 12 teams.", kb)
    assert not unsupported_terms(
        "Backend engineer working with Node.js, PostgreSQL and AWS.", kb
    )


async def test_ungrounded_llm_plan_is_replaced(kb):
    class BadClient:
        async def complete(self, prompt, schema=None, **kwargs):
            return CVPlan(
                headline="Principal Platform Architect",
                summary="Ten years of Salesforce and Kubernetes leadership.",
                experience_bullet_ids=[],
                project_ids=[],
                skill_names=[],
            )

    plan = await build_plan(make_job(), kb, client=BadClient())
    assert is_grounded(plan.headline, kb)
    assert is_grounded(plan.summary, kb)


async def test_llm_failure_falls_back_to_deterministic(kb):
    class BrokenClient:
        async def complete(self, prompt, schema=None, **kwargs):
            raise RuntimeError("upstream down")

    plan = await build_plan(make_job(), kb, client=BrokenClient())
    assert plan.experience_bullet_ids
    assert plan.summary


def test_selected_material_maps_back_to_kb(kb):
    plan = deterministic_plan(make_job(), kb)
    roles = selected_bullets(plan, kb)
    projects = selected_projects(plan, kb)

    assert roles
    assert all(bullets for _, bullets in roles)
    assert len(projects) == len(plan.project_ids)


def test_filename_is_safe_and_job_specific(kb):
    job = make_job()
    other = make_job()
    base = filename_base(job, kb.name)
    assert "/" not in base
    assert " " not in base
    assert base == base.lower()
    assert str(job.id)[:8] in base
    assert base != filename_base(other, kb.name)


def test_cv_data_contains_no_fabricated_sections(kb):
    plan = deterministic_plan(make_job(), kb)
    data = build_cv_data(make_job(), kb, plan)
    assert data.name == kb.name
    assert data.roles and data.projects
    # Certifications and achievements are empty in the KB: no section invented.
    assert data.certifications == ()


def test_renders_pdf_and_docx_with_extractable_text(kb, tmp_path):
    import pymupdf
    from docx import Document

    job = make_job()
    plan = deterministic_plan(job, kb)
    artifact = generate_cv(job, kb, plan, output_dir=tmp_path)

    assert artifact.pdf_path.exists() and artifact.pdf_path.stat().st_size > 2000
    assert artifact.docx_path.exists() and artifact.docx_path.stat().st_size > 5000
    assert artifact.pdf_path.stem == artifact.docx_path.stem

    pdf_text = "".join(page.get_text() for page in pymupdf.open(artifact.pdf_path))
    assert kb.name in pdf_text
    assert "EXPERIENCE" in pdf_text
    assert "TECHNICAL SKILLS" in pdf_text
    assert "Acme" not in pdf_text  # the CV targets the job, it does not name it

    document = Document(str(artifact.docx_path))
    assert len(document.tables) == 0  # parsers dislike tables
    docx_text = "\n".join(p.text for p in document.paragraphs)
    assert kb.name in docx_text
    assert "EXPERIENCE" in docx_text


def test_phone_is_configurable(kb, monkeypatch):
    from src.documents.generator import _contact_lines

    monkeypatch.setenv("CV_INCLUDE_PHONE", "0")
    assert kb.phone not in _contact_lines(kb)
    monkeypatch.setenv("CV_INCLUDE_PHONE", "1")
    assert kb.phone in _contact_lines(kb)
